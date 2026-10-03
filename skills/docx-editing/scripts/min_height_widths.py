# -*- coding: utf-8 -*-
"""표 높이가 가장 낮은 열 폭 — 글꼴 폭으로 칸마다 줄 수를 «재서» 고른다 (docx).

⭐ 유래 (2026-09-15, insaui4 유방촬영 답안)
  사용자가 SIGN 체크리스트 표의 «높이를 줄이려고» 손으로 열을 좁히다가 「번/호」·「Acce/ptabl/e」처럼
  토큰 중간에서 끊었다. 목적은 숫자가 아니라 높이였다. 칸마다 줄 수를 맑은 고딕 글꼴 폭으로 세어
  «두 표 합 줄 수 최소 · 토큰 중간 끊김 0» 폭을 찾았다(이 도구로 잰 값: 사람 폭 106줄·끊김 3 → 104줄·끊김 0.
  행별 줄 수는 렌더 PNG 와 대조해 맞았다).

`col_widths.py`(내용 비례 어림)와 역할이 다르다 — 그쪽은 «처음 짤 때의 기본값», 이쪽은
«이미 있는 표를 낮게 다시 짜기». 합치지 말 것.

쓰는 법
    python min_height_widths.py 문서.docx --after "SIGN Methodology Checklist 2"     # 보고만
    python min_height_widths.py 문서.docx --after "Checklist 2" --probe             # 지금 폭의 행별 줄 수
    python min_height_widths.py 문서.docx --tables 4 5 --out 새문서.docx              # 적용해 새 파일로
    python min_height_widths.py --selftest

    from min_height_widths import Meter, measure_tables, solve, apply_widths

    --after TEXT   이 글자를 포함하는 문단 «바로 다음» 표를 고른다(여러 개면 전부). 여러 표는 같은 폭을
                   공유한다 — 열 수가 같아야 한다.
    --tables N ..  본문 최상위 표 순번(0부터)
    --total-cm X   표 전체 폭(기본: 첫 표의 현재 gridCol 합)
    --lock I=CM    I 열(0부터) 폭 고정. 여러 번 줄 수 있다.
    --out PATH     적용한 문서를 여기 저장. ⛔ 입력 파일 덮어쓰기는 거부한다 — 사람 손편집이 있을 수 있다.

⚠ 함정
  1. 줄 수는 «모형»이다. Word 는 공백에서 줄을 바꾸고(한국어도 어절 단위), 칸 좌우 여백은 표 기본
     108 twips 로 본다. **적용 전에 `--probe` 로 지금 폭의 행별 줄 수를 찍어 렌더 PNG 와 두세 칸 맞춰 본다**
     (2026-09-15: 3/3 일치). 안 맞으면 여백(`--margin-tw`)이나 글꼴부터 의심한다.
  2. `audit_table_widths.py` 와 계측이 0.01in 쯤 다르다 — 거기서 BREAK 가 떠도 렌더에서 멀쩡하면
     이 모형이 맞다(2026-09-15 「번호」: 감사기 0.25in 필요 / 칸 0.24in / 렌더 정상). 판정은 렌더로.
  3. 가로 병합(gridSpan)은 합친 폭으로 잰다. 세로 병합의 이어지는 칸은 글자가 없어 줄 수에 안 들어간다.
  4. 폭만 바꾼다. 글자 크기·여백은 건드리지 않는다 — 높이를 줄이려고 글씨를 줄이지 않는다.
  4-b. 글자 크기는 «run» 의 rPr 에서 읽는다. 첫 판(즉석 스크립트)은 `.//rPr` 로 문단 기호의 rPr 을 먼저 집어
     9pt 칸을 10pt 로 재어 두 표 높이를 4줄 부풀렸다(110 vs 실제 106). 상대 비교는 맞아도 절대값이 틀린다.
  5. 탐색은 국소 탐색(열 쌍 사이로 폭을 옮기며 줄 수가 줄면 받아들임)을 출발점 여러 개로 돌린다.
     열이 많으면(>6) 전역 최적을 보장하지 않는다 — 결과 줄 수를 지금 폭과 나란히 보고하므로 비교해서 쓴다.
"""
import argparse
import math
import os
import sys

from docx import Document
from docx.oxml.ns import qn

TW_PER_CM = 566.93
DEFAULT_MARGIN_TW = 108
FONT_DIR = os.path.join(os.environ.get('WINDIR', r'C:\Windows'), 'Fonts')
DEFAULT_FONTS = {False: 'malgun.ttf', True: 'malgunbd.ttf'}


def W(tag):
    return qn('w:' + tag)


# ── 1. 계측 ──────────────────────────────────────────────────────────────
class Meter:
    """글꼴 폭으로 한 칸의 줄 수를 센다."""

    def __init__(self, fonts=None, margin_tw=DEFAULT_MARGIN_TW):
        from PIL import ImageFont
        fonts = fonts or DEFAULT_FONTS
        self._font = {}
        for bold, name in fonts.items():
            path = name if os.path.isabs(name) else os.path.join(FONT_DIR, name)
            if not os.path.exists(path):
                raise FileNotFoundError('글꼴이 없다: %s — fonts= 로 경로를 줄 것' % path)
            self._font[bold] = ImageFont.truetype(path, 1000)
        self._cache = {}
        self.margin_pt = 2 * margin_tw / 20.0

    def text_pt(self, s, bold, size_pt):
        key = (s, bold)
        if key not in self._cache:
            self._cache[key] = self._font[bold].getlength(s) / 1000.0
        return self._cache[key] * size_pt

    def lines(self, text, width_tw, bold, size_pt):
        """(줄 수, 토큰 중간 끊김 여부)."""
        avail = width_tw / 20.0 - self.margin_pt
        if avail <= 0:
            return 999, True
        sp = self.text_pt(' ', bold, size_pt)
        total, broke = 0, False
        for para in text.split('\n'):
            n, cur = 1, 0.0
            for word in para.split(' '):
                if not word:
                    continue
                wl = self.text_pt(word, bold, size_pt)
                if wl > avail:
                    broke = True
                    k = math.ceil(wl / avail)
                    n += (k - 1) + (1 if cur > 0 else 0)
                    cur = wl - (k - 1) * avail
                    continue
                if cur == 0:
                    cur = wl
                elif cur + sp + wl <= avail:
                    cur += sp + wl
                else:
                    n += 1
                    cur = wl
            total += n
        return total, broke


# ── 2. 표 읽기 ────────────────────────────────────────────────────────────
def _default_size_pt(doc):
    sz = doc.styles.element.find('.//' + W('docDefaults') + '//' + W('sz'))
    return int(sz.get(W('val'))) / 2.0 if sz is not None else 10.0


def measure_tables(doc, tbls):
    """표들을 행 목록으로: 각 행 = [(글자, 시작열, 칸수, 굵게, 크기pt)]."""
    base = _default_size_pt(doc)
    rows = []
    for tbl in tbls:
        for tr in tbl.findall(W('tr')):
            cells, col = [], 0
            for tc in tr.findall(W('tc')):
                tcpr = tc.find(W('tcPr'))
                g = tcpr.find(W('gridSpan')) if tcpr is not None else None
                span = int(g.get(W('val'))) if g is not None else 1
                paras = []
                for p in tc.findall(W('p')):
                    s = ''
                    for r in p.iter(W('r')):
                        for ch in r:
                            if ch.tag == W('t'):
                                s += ch.text or ''
                            elif ch.tag == W('br'):
                                s += '\n'
                    paras.append(s)
                rpr = tc.find('.//' + W('r') + '/' + W('rPr'))
                b = rpr.find(W('b')) if rpr is not None else None
                bold = b is not None and b.get(W('val')) not in ('0', 'false')
                sz = rpr.find(W('sz')) if rpr is not None else None
                size = int(sz.get(W('val'))) / 2.0 if sz is not None else base
                cells.append(('\n'.join(paras), col, span, bold, size))
                col += span
            rows.append(cells)
    return rows


def current_widths(tbl):
    return [int(g.get(W('w'))) for g in tbl.find(W('tblGrid')).findall(W('gridCol'))]


# ── 3. 평가 · 탐색 ────────────────────────────────────────────────────────
def evaluate(rows, widths, meter):
    """(표 높이 = 행마다 최대 줄 수의 합, 토큰 중간 끊김 목록)."""
    total, breaks = 0, []
    for r, cells in enumerate(rows):
        m = 1
        for text, c0, span, bold, size in cells:
            if not text.strip():
                continue
            n, b = meter.lines(text, sum(widths[c0:c0 + span]), bold, size)
            if b:
                breaks.append((r, c0, text[:16]))
            m = max(m, n)
        total += m
    return total, breaks


def _min_widths(rows, ncols, meter, locks, total_tw):
    mins = []
    for c in range(ncols):
        if c in locks:
            mins.append(locks[c])
            continue
        cells = [x for row in rows for x in row if x[1] == c and x[2] == 1 and x[0].strip()]
        w = 200
        while w < total_tw and any(meter.lines(t, w, b, s)[1] for t, _, _, b, s in cells):
            w += 14
        mins.append(w)
    return mins


def solve(rows, ncols, total_tw, meter, locks=None, ref=None, steps=(283, 142, 57, 28)):
    """높이 최소·끊김 0 폭(twips 목록)과 (줄 수, 끊김)을 돌려준다.

    ref(사람이 준 폭 등)가 있으면 같은 높이일 때 그쪽에 가까운 폭을 고른다.
    """
    locks = locks or {}
    mins = _min_widths(rows, ncols, meter, locks, total_tw)
    if sum(mins) > total_tw:
        raise ValueError('토큰이 안 끊기는 최소 폭 합 %.2fcm 가 표 폭 %.2fcm 보다 크다 — 표를 넓히거나 글을 줄일 것'
                         % (sum(mins) / TW_PER_CM, total_tw / TW_PER_CM))
    free = [c for c in range(ncols) if c not in locks]

    def fill(base_extra):
        ws = list(mins)
        spare = total_tw - sum(ws)
        wsum = sum(base_extra[c] for c in free) or 1
        for c in free:
            ws[c] += int(spare * base_extra[c] / wsum)
        ws[free[-1]] += total_tw - sum(ws)
        return ws

    content = [0.0] * ncols
    for row in rows:
        for t, c0, span, b, s in row:
            if span == 1:
                content[c0] += meter.text_pt(t.replace('\n', ' '), b, s)
    starts = [fill(content), fill([1.0] * ncols)]
    if ref and len(ref) == ncols:
        scale = total_tw / float(sum(ref))
        r2 = [max(mins[c], int(ref[c] * scale)) if c not in locks else locks[c] for c in range(ncols)]
        r2[free[-1]] += total_tw - sum(r2)
        if r2[free[-1]] >= mins[free[-1]]:
            starts.append(r2)

    def key(ws):
        lines, br = evaluate(rows, ws, meter)
        dist = sum(abs(a - b) for a, b in zip(ws, ref)) if ref and len(ref) == ncols else 0
        return (len(br), lines, dist)

    best_ws, best_k = None, None
    for ws in starts:
        cur, ck = list(ws), key(ws)
        for step in steps:
            improved = True
            while improved:
                improved = False
                cand_best, cand_k = None, ck
                for i in free:
                    for j in free:
                        if i == j or cur[i] - step < mins[i]:
                            continue
                        cand = list(cur)
                        cand[i] -= step
                        cand[j] += step
                        k = key(cand)
                        if k < cand_k:
                            cand_best, cand_k = cand, k
                if cand_best is not None:
                    cur, ck, improved = cand_best, cand_k, True
        if best_k is None or ck < best_k:
            best_ws, best_k = cur, ck
    lines, br = evaluate(rows, best_ws, meter)
    return best_ws, lines, br


def apply_widths(tbl, widths):
    """gridCol 과 각 칸 tcW(가로 병합은 합)를 widths(twips)로 맞춘다."""
    grid = tbl.find(W('tblGrid'))
    for g, w in zip(grid.findall(W('gridCol')), widths):
        g.set(W('w'), str(int(w)))
    for tr in tbl.findall(W('tr')):
        col = 0
        for tc in tr.findall(W('tc')):
            tcpr = tc.find(W('tcPr'))
            g = tcpr.find(W('gridSpan')) if tcpr is not None else None
            span = int(g.get(W('val'))) if g is not None else 1
            tcw = tcpr.find(W('tcW')) if tcpr is not None else None
            if tcw is not None:
                tcw.set(W('w'), str(int(sum(widths[col:col + span]))))
                tcw.set(W('type'), 'dxa')
            col += span


# ── 4. 표 고르기 · CLI ────────────────────────────────────────────────────
def select_tables(doc, after=None, indices=None):
    body = doc.element.body
    kids = list(body.iterchildren())
    tops = [k for k in kids if k.tag == W('tbl')]
    if indices:
        return [tops[i] for i in indices]
    out = []
    for i, el in enumerate(kids):
        if el.tag == W('p') and after in ''.join(t.text or '' for t in el.iter(W('t'))):
            for nxt in kids[i + 1:]:
                if nxt.tag == W('tbl'):
                    out.append(nxt)
                    break
                if nxt.tag == W('p') and ''.join(t.text or '' for t in nxt.iter(W('t'))).strip() \
                        and after in ''.join(t.text or '' for t in nxt.iter(W('t'))):
                    break
    return out


def _cm(ws):
    return '[' + ' / '.join('%.2f' % (w / TW_PER_CM) for w in ws) + '] cm'


def main(argv):
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('docx', nargs='?')
    ap.add_argument('--after')
    ap.add_argument('--tables', type=int, nargs='+')
    ap.add_argument('--total-cm', type=float)
    ap.add_argument('--lock', action='append', default=[])
    ap.add_argument('--margin-tw', type=int, default=DEFAULT_MARGIN_TW)
    ap.add_argument('--probe', action='store_true')
    ap.add_argument('--out')
    ap.add_argument('--selftest', action='store_true')
    a = ap.parse_args(argv)
    if a.selftest:
        return selftest()
    if not a.docx or not (a.after or a.tables):
        ap.error('docx 와 --after 또는 --tables 가 필요하다')
    doc = Document(a.docx)
    tbls = select_tables(doc, a.after, a.tables)
    if not tbls:
        sys.exit('고른 표가 없다')
    ncols = {len(current_widths(t)) for t in tbls}
    if len(ncols) != 1:
        sys.exit('열 수가 다른 표를 한 폭으로 풀 수 없다: %s' % sorted(ncols))
    ncols = ncols.pop()
    meter = Meter(margin_tw=a.margin_tw)
    rows = measure_tables(doc, tbls)
    cur = current_widths(tbls[0])
    total = int(round(a.total_cm * TW_PER_CM)) if a.total_cm else sum(cur)
    locks = {}
    for spec in a.lock:
        i, cmv = spec.split('=')
        locks[int(i)] = int(round(float(cmv) * TW_PER_CM))
    lines0, br0 = evaluate(rows, cur, meter)
    print('표 %d개 · 열 %d · 행 %d' % (len(tbls), ncols, len(rows)))
    print('지금 폭   %s → 높이 %d줄 · 토큰 중간 끊김 %d' % (_cm(cur), lines0, len(br0)))
    for r, c, t in br0:
        print('    끊김: 행 %d 열 %d 「%s」' % (r, c, t))
    if a.probe:
        for r, cells in enumerate(rows):
            n = max([1] + [meter.lines(t, sum(cur[c0:c0 + sp]), b, s)[0]
                           for t, c0, sp, b, s in cells if t.strip()])
            print('    행 %2d  %d줄  %s' % (r, n, (cells[0][0] if cells else '')[:24].replace('\n', ' ')))
    ws, lines, br = solve(rows, ncols, total, meter, locks=locks, ref=cur)
    print('고른 폭   %s → 높이 %d줄 · 토큰 중간 끊김 %d' % (_cm(ws), lines, len(br)))
    if a.out:
        if os.path.abspath(a.out) == os.path.abspath(a.docx):
            sys.exit('⛔ 입력 파일 덮어쓰기는 하지 않는다 — 다른 --out 을 줄 것')
        for t in tbls:
            apply_widths(t, ws)
        doc.save(a.out)
        print('저장:', a.out, '— 렌더해서 눈으로 볼 것')
    return 0


# ── 5. 자기 시험 ──────────────────────────────────────────────────────────
def selftest():
    sys.stdout.reconfigure(encoding='utf-8')
    ok = True
    try:
        meter = Meter()
    except FileNotFoundError as e:
        print('건너뜀:', e)
        return 0
    rows = [
        [('번호', 0, 1, True, 9), ('항목', 1, 1, True, 9), ('판정', 2, 1, True, 9), ('근거', 3, 1, True, 9)],
        [('1.10', 0, 1, False, 9),
         ('Where the study is carried out at more than one site, results are comparable for all sites.', 1, 1, False, 9),
         ('Acceptable (+)', 2, 1, False, 9),
         ('스웨덴 4개 검진기관, 단일 벤더 장비로 표준화했고 판독의는 배정을 알았다', 3, 1, False, 9)],
        [('SECTION 2: OVERALL ASSESSMENT OF THE STUDY', 0, 4, True, 9)],
        [('종합', 0, 2, False, 9), ('1+', 2, 1, False, 9), ('설계 1 × 2.1의 + = 1+', 3, 1, False, 9)],
    ]
    total = int(16 * TW_PER_CM)
    narrow = [int(1.0 * TW_PER_CM), int(6.0 * TW_PER_CM), int(1.25 * TW_PER_CM), total]
    narrow[3] = total - sum(narrow[:3])
    n0, b0 = evaluate(rows, narrow, meter)
    ws, n1, b1 = solve(rows, 4, total, meter, ref=narrow)
    for label, cond in (
            ('[음성] 좁은 판정 열에서 「Acceptable」 끊김을 잡는다', any('Acceptable' in t for _, _, t in b0)),
            ('[양성] 푼 폭에는 끊김이 없다', not b1),
            ('[양성] 푼 폭의 합이 표 폭과 같다', sum(ws) == total),
            ('[양성] 높이가 좁은 폭보다 높지 않다', n1 <= n0)):
        print('%-40s %s' % (label, '✓' if cond else '🔴'))
        ok &= bool(cond)
    ws2, _, _ = solve(rows, 4, total, meter, locks={0: int(1.2 * TW_PER_CM)})
    cond = ws2[0] == int(1.2 * TW_PER_CM)
    print('%-40s %s' % ('[양성] --lock 열은 움직이지 않는다', '✓' if cond else '🔴'))
    ok &= cond
    try:
        solve(rows, 4, int(3 * TW_PER_CM), meter)
        print('%-40s 🔴' % '[음성] 표가 너무 좁으면 멈춘다')
        ok = False
    except ValueError:
        print('%-40s ✓' % '[음성] 표가 너무 좁으면 멈춘다')
    print('-' * 50)
    print('자기 시험:', '통과' if ok else '🔴 실패')
    return 0 if ok else 1


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    sys.exit(main(sys.argv[1:]))
